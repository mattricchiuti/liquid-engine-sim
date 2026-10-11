import csv
import itertools
import math
import os
import time
import webbrowser
from pathlib import Path

import numpy as np                                # arrays and math on whole time histories
import openpyxl                                   # reading input.xlsx
import plotly.graph_objects as go                 # interactive plots (web pages)
from plotly.subplots import make_subplots
from CoolProp.CoolProp import PropsSI             # N2O properties: PropsSI(output, in1, v1, in2, v2, "NitrousOxide")
from rocketcea.cea_obj_w_units import CEA_Obj     # NASA CEA (SI-units version): c*, Tc, gamma, MW vs. Pc and O/F

INPUT_FILE = "input.xlsx"

# HalfCat switches the N2O tank to gas-phase blowdown below this fraction of the starting ox mass
GAS_PHASE_FRACTION = 0.10

# CoolProp has no N2O viscosity model; approximate value for the vapor (only affects line friction)
N2O_VAPOR_VISCOSITY = 1.5e-5  # [Pa s]

# O/F is clamped to this range before calling CEA (it fails at extreme mixtures, e.g. zero fuel flow)
OF_MIN, OF_MAX = 0.5, 10.0

# CEA lookup table grid: chamber pressure spaced evenly in ln(Pc) (properties vary slowly with Pc);
# O/F finer below 2, where properties change fastest (fuel-rich tail)
CEA_PC_GRID = np.geomspace(5e4, 8e6, 20)  # [Pa]
CEA_OF_GRID = np.concatenate([np.arange(OF_MIN, 2.0, 0.05), np.arange(2.0, OF_MAX + 1e-9, 0.1)])

G0 = 9.80665  # standard gravity [m/s^2], for Isp

# Rule of thumb: nozzle flow may separate from the wall when exit pressure < this fraction of ambient
SEPARATION_RATIO = 0.4

# Startup ramp (after HalfCat): flow is established before ignition, then chamber pressure rises
# linearly from ambient to its balanced value over STARTUP_TIME, stepped at STARTUP_DT
STARTUP_TIME = 0.010  # [s]
STARTUP_DT = 0.001    # [s]

# Column letters for each input panel on the sheet: parameter name, then (value, unit) pairs
PANELS = [
    {"name": "C", "imperial": ("D", "E"), "metric": ("F", "G")},
    {"name": "L", "imperial": ("M", "N"), "metric": ("O", "P")},
]

# Unit written in the spreadsheet -> function converting a value in that unit to SI
TO_SI = {
    "": lambda x: x,                          # unitless (counts, Cd, Cv, efficiencies)
    "s": lambda x: x,
    "1/s": lambda x: x,
    "m": lambda x: x,
    "Pa": lambda x: x,
    "kg/m3": lambda x: x,
    "in": lambda x: x * 0.0254,               # -> m
    "ft": lambda x: x * 0.3048,               # -> m
    "psi": lambda x: x * 6894.757,            # -> Pa
    "F": lambda x: (x - 32) / 1.8 + 273.15,   # -> K
    "deg": math.radians,                      # -> rad
    "lbm/ft3": lambda x: x * 16.01846,        # -> kg/m^3
    "cP": lambda x: x * 0.001,                # -> Pa*s
}


# ---------- inputs ----------
def load_inputs(path=INPUT_FILE):
    """Read the 'input' sheet.

    Returns (inputs, sweeps):
      inputs: {parameter name: value in SI}
      sweeps: list of active sweeps, each {"parameter": name, "values": [values in SI],
              "as_entered": [values in spreadsheet units], "unit": spreadsheet unit}

    For each parameter, uses whichever column (Imperial or Metric) holds a
    typed value rather than a formula, preferring Imperial if both do; if both
    are formulas, uses the calculated value. Converts using the unit written
    next to the value. An unrecognized unit is an error.
    """
    # data_only=True gives the values formulas last calculated to; the second
    # copy keeps the formula text, so typed values can be told apart from formulas
    ws = openpyxl.load_workbook(path, data_only=True)["input"]
    ws_formulas = openpyxl.load_workbook(path)["input"]
    inputs = {"Oxidizer": "N2O"}  # the tank model only handles self-pressurizing N2O
    units = {}  # unit each parameter was read in, for converting sweep values

    def is_typed_number(cell_ref):
        return isinstance(ws_formulas[cell_ref].value, (int, float))

    # Fuel name: the dropdown to the right of the "Propellant Properties" heading
    for row in ws.iter_rows():
        for cell in row:
            if cell.value == "Propellant Properties":
                inputs["Fuel"] = ws.cell(cell.row, cell.column + 1).value
    if "Fuel" not in inputs:
        raise ValueError(f"No 'Propellant Properties' heading (with the fuel dropdown beside it) in {path}")

    # Sweep block: rows under the "Parameter | min | max | step | 0/1" header
    sweep_rows = []
    for row in range(1, ws.max_row + 1):
        if ws[f"C{row}"].value == "Parameter" and ws[f"D{row}"].value == "min":
            r = row + 1
            while isinstance(ws[f"G{r}"].value, (int, float)):  # each sweep row has a 0/1 toggle
                sweep_rows.append(r)
                r += 1

    unknown_units = []
    for col in PANELS:
        for row in range(1, ws.max_row + 1):
            if col["name"] == "C" and row in sweep_rows:
                continue
            label = ws[f"{col['name']}{row}"].value
            columns = (col["imperial"], col["metric"])
            # A typed value wins; if both columns are formulas (e.g. Oxidizer Length), use the calculated one
            source = next((c for c in columns if is_typed_number(f"{c[0]}{row}")), None) \
                or next((c for c in columns if isinstance(ws[f"{c[0]}{row}"].value, (int, float))), None)
            if not isinstance(label, str) or source is None:
                continue  # section titles, column headers, blank rows
            value = ws[f"{source[0]}{row}"].value
            unit = (ws[f"{source[1]}{row}"].value or "").strip()
            label = label.strip()
            if label in inputs:
                raise ValueError(f"Duplicate parameter name {label!r} in {path}")
            if unit not in TO_SI:
                unknown_units.append(f"{label} ({unit!r}, row {row})")
                continue
            inputs[label] = TO_SI[unit](value)
            units[label] = unit

    if unknown_units:
        raise ValueError(f"No SI conversion for the unit of: " + "; ".join(unknown_units))

    sweeps = []
    for row in sweep_rows:
        if ws[f"G{row}"].value != 1:
            continue  # sweep turned off
        name = (ws[f"C{row}"].value or "").strip()
        lo, hi, step = (ws[f"{c}{row}"].value for c in "DEF")
        if name not in inputs:
            raise ValueError(f"Sweep in row {row} is on but its parameter {name!r} isn't an input")
        if not all(isinstance(v, (int, float)) for v in (lo, hi, step)) or step <= 0 or hi < lo:
            raise ValueError(f"Sweep in row {row} ({name}) needs numbers with min <= max and step > 0")
        n = round((hi - lo) / step) + 1  # both ends included
        as_entered = [lo + i * step for i in range(n)]
        sweeps.append({
            "parameter": name,
            "values": [TO_SI[units[name]](v) for v in as_entered],  # SI, for the sim
            "as_entered": as_entered,                               # spreadsheet units, for labels
            "unit": units[name],
        })

    return inputs, sweeps


# math functions

def area(diameter):
    return math.pi / 4 * (diameter ** 2)

def bisect(func, lo, hi, tol=1.0, max_iter=100):
    """Find x in [lo, hi] where func(x) = 0; func(lo) and func(hi) must have opposite signs."""
    f_lo = func(lo)
    if f_lo * func(hi) > 0:
        raise ValueError(f"bisect: no sign change between {lo} and {hi}")
    for _ in range(max_iter):
        mid = (lo + hi) / 2
        f_mid = func(mid)
        if f_mid * f_lo > 0:        # same sign as the low end, so the root is in the upper half
            lo, f_lo = mid, f_mid
        else:
            hi = mid
        if hi - lo < tol:
            break
    return (lo + hi) / 2


# engine calcs

def initial_conditions(inputs):

    tank_ID = inputs["Tank OD"] - 2 * inputs["Tank Wall Thickness"]
    tank_area = math.pi / 4 * tank_ID ** 2

    T_ox = inputs["Nitrous Temperature"]

    rho_ox = PropsSI("D", "T", T_ox, "Q", 0, inputs["Oxidizer"])
    V_ox = inputs["Oxidizer Length"] * tank_area
    m_ox0 = rho_ox * V_ox

    rho_fuel = inputs["Fuel Density"]
    V_fuel = inputs["Fuel Length"] * tank_area
    m_fuel0 = rho_fuel * V_fuel

    P_ox0 = PropsSI("P", "T", T_ox, "Q", 0, inputs["Oxidizer"])
    P_fuel0 = P_ox0 - inputs["Fuel Piston Loss"]

    A_t = area(inputs["Throat Diameter"])
    A_e = area(inputs["Exit Diameter"])
    A_oxfeed = area(inputs["Oxidizer Line Diameter"])
    A_fuelfeed = area(inputs["Fuel Line Diameter"])

    CdA_ox = inputs["No. Oxidizer Holes"] * area(inputs["Oxidizer Hole Diameter"]) * inputs["Oxidizer Cd"]
    CdA_fuel = inputs["No. Fuel Holes"] * area(inputs["Fuel Hole Diameter"]) * inputs["Fuel Cd"]

    return {
        "T_ox": T_ox,
        "rho_ox": rho_ox,
        "rho_fuel": rho_fuel,
        "V_ox": V_ox,
        "V_fuel": V_fuel,
        "m_ox0": m_ox0,
        "m_fuel0": m_fuel0,
        "P_ox0": P_ox0,
        "P_fuel0": P_fuel0,
        "A_t": A_t,
        "A_e": A_e,
        "eps": A_e / A_t,  # nozzle area (expansion) ratio
        "A_oxfeed": A_oxfeed,
        "A_fuelfeed": A_fuelfeed,
        "CdA_ox": CdA_ox,
        "CdA_fuel": CdA_fuel,
    }

def tank_pressures(m_ox, t, t_switch, ic, inputs):
    """Tank pressures and N2O density for the oxidizer mass left (HalfCat's empirical blowdown).

    m_ox:     oxidizer mass left now [kg]
    t:        time now [s]
    t_switch: time the tank switched to the gas phase [s], or None while still liquid
    """
    L = inputs["Decay Constant (liquid)"]
    k = inputs["Decay Constant (vapor)"]
    fraction_left = m_ox / ic["m_ox0"]

    if fraction_left > GAS_PHASE_FRACTION:
        # Liquid phase: straight line from P0 (full) down to L*P0 (empty)
        phase = "liquid"
        P_ox = ic["P_ox0"] * (fraction_left * (1 - L) + L)
        rho_ox = PropsSI("D", "P", P_ox, "Q", 0, inputs["Oxidizer"])  # saturated liquid
    else:
        # Gas phase: exponential decay from the pressure where the liquid phase ended
        phase = "gas"
        if t_switch is None:
            t_switch = t  # first step in the gas phase
        P_switch = ic["P_ox0"] * (GAS_PHASE_FRACTION * (1 - L) + L)
        P_ox = P_switch * math.exp(-k * (t - t_switch))
        rho_ox = PropsSI("D", "P", P_ox, "Q", 1, inputs["Oxidizer"])  # saturated vapor

    P_fuel = P_ox - inputs["Fuel Piston Loss"]

    return {
        "phase": phase,
        "t_switch": t_switch,  # the loop keeps this and passes it back in next step
        "P_ox": P_ox,
        "P_fuel": P_fuel,
        "rho_ox": rho_ox,
    }


def friction_factor(Re, rel_roughness):
    """Darcy friction factor: laminar below Re = 2300, Haaland equation above.

    rel_roughness: line roughness / line diameter
    """
    if Re < 2300:
        return 64 / Re
    return (-1.8 * math.log10((rel_roughness / 3.7) ** 1.11 + 6.9 / Re)) ** -2


def fuel_flow(P_fuel, Pc, ic, inputs):
    """Fuel mass flow [kg/s] from the fuel tank to the chamber (single-phase incompressible).

    Line + fittings and the injector are resistances in series:
        P_fuel - Pc = mdot^2 / (2 rho) * (R_feed + R_inj)
    """
    if P_fuel <= Pc:
        return 0.0  # no pressure drop, no flow

    rho = ic["rho_fuel"]
    D = inputs["Fuel Line Diameter"]
    L = inputs["Fuel Line Length"]
    A_line = ic["A_fuelfeed"]
    R_inj = 1 / ic["CdA_fuel"] ** 2

    f = 0.02  # starting guess; refined below from the Reynolds number
    for _ in range(5):
        R_feed = (f * L / D + inputs["Fuel Sum K"]) / A_line ** 2
        mdot = math.sqrt(2 * rho * (P_fuel - Pc) / (R_feed + R_inj))
        Re = 4 * mdot / (math.pi * D * inputs["Fuel Dynamic Viscosity"])
        f = friction_factor(Re, inputs["Fuel Line Roughness"] / D)

    return mdot


def ox_upstream(P_ox, inputs, phase="liquid", n_grid=60):
    """Tank-side N2O state for the current tank pressure and phase ("liquid" or "gas").

    Called once per time step: everything here depends only on tank pressure, so the
    bisections inside ox_flow can reuse it instead of calling CoolProp thousands of times.
    Liquid: saturated liquid plus the HEM flux curve for the NHNE injector model.
    Gas: saturated vapor plus its isentropic exponent for compressible injector flow.
    """
    ox = inputs["Oxidizer"]

    if phase == "gas":
        rho_v = PropsSI("D", "P", P_ox, "Q", 1, ox)
        a = PropsSI("A", "P", P_ox, "Q", 1, ox)  # speed of sound
        return {
            "phase": "gas",
            "T1": PropsSI("T", "P", P_ox, "Q", 1, ox),
            "rho1": rho_v,
            "mu1": N2O_VAPOR_VISCOSITY,
            "P_tank": P_ox,
            # Isentropic exponent from the speed of sound, k = rho a^2 / P. Saturated N2O vapor is far
            # from ideal (cp/cv ~1.9 vs k ~1.2), and this k keeps ideal-gas nozzle formulas within ~5%
            # of a full real-gas calculation.
            "k": rho_v * a ** 2 / P_ox,
        }

    # Saturated liquid leaving the tank (Q = 0)
    T1 = PropsSI("T", "P", P_ox, "Q", 0, ox)
    rho_l = PropsSI("D", "P", P_ox, "Q", 0, ox)
    h1 = PropsSI("H", "P", P_ox, "Q", 0, ox)
    s1 = PropsSI("S", "P", P_ox, "Q", 0, ox)

    # HEM: expand at constant entropy from (h1, s1) down to each downstream pressure P2;
    # the enthalpy drop becomes velocity, v = sqrt(2 * (h1 - h2)), and flux = rho2 * v
    P_grid = np.linspace(inputs["Ambient Pressure"], P_ox, n_grid)
    h2 = PropsSI("H", "P", P_grid, "S", s1, ox)
    rho2 = PropsSI("D", "P", P_grid, "S", s1, ox)
    G = rho2 * np.sqrt(np.maximum(2 * (h1 - h2), 0))  # max(..., 0) guards rounding at P2 = P_ox

    # The liquid is assumed to stay liquid through the feed line ("frozen"): friction drops its
    # pressure but keeps h1 and s1, so this curve applies at the manifold too. Choking (taking the
    # max of the curve between Pc and the manifold pressure) happens in nhne_flux, where P_man is known.
    # The curve has a single hump, so that max is either the peak or one of the two ends: store the peak.
    i_peak = int(np.argmax(G))

    return {
        "phase": "liquid",
        "T1": T1,
        "rho1": rho_l,
        "mu1": inputs["Oxidizer Dynamic Viscosity"],  # CoolProp has no N2O viscosity model
        "h1": h1,
        "s1": s1,
        "P_v": P_ox,  # saturated tank: vapor pressure = tank pressure
        "P_grid": P_grid,
        "G_hem": G,  # raw (unchoked) HEM flux at each pressure in P_grid
        "P_peak": P_grid[i_peak],  # where the HEM flux peaks (chokes)
        "G_peak": G[i_peak],
    }


def nhne_flux(P1, P2, up):
    """NHNE (Dyer) mass flux, G [kg/m^2/s] through the injector from manifold P1 to chamber P2.

    up: tank-side N2O state from ox_upstream. Multiply the result by CdA to get mass flow.
    """
    if P1 <= P2:
        return 0.0  # no forward pressure drop, no flow

    # SPI: liquid stays liquid through the hole (upper limit)
    G_spi = math.sqrt(2 * up["rho1"] * (P1 - P2))

    # HEM: liquid boils instantly (lower limit), choked: the largest flux on the curve between
    # the chamber and the manifold. The curve has one hump, so that's the peak if it lies in
    # between (choked), otherwise the larger of the two ends.
    if P2 <= up["P_peak"] <= P1:
        G_hem = up["G_peak"]
    else:
        G_hem = max(np.interp(P2, up["P_grid"], up["G_hem"]), np.interp(P1, up["P_grid"], up["G_hem"]))

    # K (Dyer's kappa): bubble growth time vs. time in the hole; large -> SPI-like, small -> HEM-like
    K = math.sqrt((P1 - P2) / (up["P_v"] - P2))

    # Blend: weights add to 1, so the result always lies between HEM and SPI
    return K / (1 + K) * G_spi + 1 / (1 + K) * G_hem


def gas_flux(P1, P2, up):
    """Mass flux, G [kg/m^2/s] of N2O vapor through the injector from manifold P1 to chamber P2.

    Compressible (ideal-gas) orifice flow using the real-gas exponent k: choked or subsonic.
    """
    if P1 <= P2:
        return 0.0

    k = up["k"]
    rho_in = up["rho1"] * P1 / up["P_tank"]  # throttled through the line at ~constant T: density ~ pressure
    r = P2 / P1
    r_choke = (2 / (k + 1)) ** (k / (k - 1))  # below this pressure ratio the hole is choked
    if r <= r_choke:
        return math.sqrt(k * P1 * rho_in * (2 / (k + 1)) ** ((k + 1) / (k - 1)))
    return math.sqrt(2 * rho_in * P1 * k / (k - 1) * (r ** (2 / k) - r ** ((k + 1) / k)))

               
def ox_flow(P_ox, Pc, up, ic, inputs):
    """Ox mass flow [kg/s] and manifold pressure [Pa] from the ox tank to the chamber.

    Feed line (friction + fittings) in series with the injector: NHNE while the tank holds
    liquid, compressible gas flow once it's vapor. Neither scales as mdot^2, so instead of a
    closed form, bisect on the manifold pressure P_man until the line delivers exactly what
    the injector passes.
    up: tank-side N2O state from ox_upstream (computed once per time step).
    """
    if P_ox <= Pc:
        return {"mdot": 0.0, "P_man": P_ox}  # no pressure drop, no flow

    injector_flux = nhne_flux if up["phase"] == "liquid" else gas_flux
    rho = up["rho1"]
    mu = up["mu1"]
    D = inputs["Oxidizer Line Diameter"]
    L = inputs["Oxidizer Line Length"]
    rel_roughness = inputs["Oxidizer Line Roughness"] / D
    A_line = ic["A_oxfeed"]
    CdA = ic["CdA_ox"]

    f = 0.02  # starting guess; refined below from the Reynolds number
    for _ in range(2):  # f settles after one update (N2O's high Reynolds number)
        R_feed = (f * L / D + inputs["Oxidizer Sum K"]) / A_line ** 2

        def mismatch(P_man):
            """Feed line flow minus injector flow, for a guessed manifold pressure."""
            line = math.sqrt(2 * rho * (P_ox - P_man) / R_feed)
            injector = CdA * injector_flux(P_man, Pc, up)
            return line - injector

        # At P_man = Pc the line flows most and the injector 0 (positive); at P_man = P_ox
        # the line flows 0 and the injector flows (negative), so the root is in between
        P_man = bisect(mismatch, Pc, P_ox, tol=10.0)
        mdot = CdA * injector_flux(P_man, Pc, up)

        Re = 4 * mdot / (math.pi * D * mu)
        f = friction_factor(Re, rel_roughness)

    return {"mdot": mdot, "P_man": P_man}


_CEA_TABLES = {}  # built tables, by (oxidizer, fuel), reused across runs (e.g. every case of a sweep)


def cea_table(inputs):
    """Table of ideal chamber properties (c*, Tc, MW, gamma) over a grid of Pc and O/F, from NASA CEA.

    Built once per propellant pair (~1 s) and reused; combustion() interpolates in it, which is
    ~40x faster than calling CEA directly. Chamber properties don't depend on the nozzle, so the
    table doesn't either.
    """
    key = (inputs["Oxidizer"], inputs["Fuel"])
    if key not in _CEA_TABLES:
        cea = CEA_Obj(oxName=key[0], fuelName=key[1], pressure_units="Pa", cstar_units="m/s",
                      temperature_units="K", isp_units="sec")
        lnP = np.log(CEA_PC_GRID)
        values = np.array([[cea.get_IvacCstrTc_ChmMwGam(Pc=Pc, MR=OF, eps=4.0)[1:]  # eps doesn't affect these
                            for OF in CEA_OF_GRID] for Pc in CEA_PC_GRID])          # shape (nPc, nOF, 4)
        _CEA_TABLES[key] = {"lnP": lnP, "OF": CEA_OF_GRID, "values": values}
    return _CEA_TABLES[key]


def combustion(Pc, OF, table, inputs):
    """Hot-gas properties in the chamber for chamber pressure Pc [Pa] and mixture ratio OF.

    Interpolates the CEA table (bilinear in ln(Pc) and O/F). CEA gives ideal (complete
    combustion) values; the c* efficiency scales c* by eta and chamber temperature by eta^2,
    which keeps them consistent since c* ~ sqrt(R * Tc).
    """
    lnP, OFs, v = table["lnP"], table["OF"], table["values"]
    x = min(max(math.log(Pc), lnP[0]), lnP[-1])  # clamp to the table
    y = min(max(OF, OF_MIN), OF_MAX)

    # Grid cell containing (x, y), and how far across it the point sits (0..1 each way)
    i = min(max(int(np.searchsorted(lnP, x)) - 1, 0), len(lnP) - 2)
    j = min(max(int(np.searchsorted(OFs, y)) - 1, 0), len(OFs) - 2)
    fx = (x - lnP[i]) / (lnP[i + 1] - lnP[i])
    fy = (y - OFs[j]) / (OFs[j + 1] - OFs[j])
    cstar, Tc, MW, gamma = ((1 - fx) * (1 - fy) * v[i, j] + fx * (1 - fy) * v[i + 1, j]
                            + (1 - fx) * fy * v[i, j + 1] + fx * fy * v[i + 1, j + 1])

    eta = inputs["C* Efficiency"]
    return {
        "cstar": eta * cstar,        # [m/s]
        "Tc": eta ** 2 * Tc,         # [K]
        "MW": MW,                    # [kg/kmol]
        "gamma": gamma,              # chamber value
        "R": 8314.46 / MW,           # specific gas constant [J/(kg K)]
    }


def chamber_state(Pc, P_ox, P_fuel, up, cea, ic, inputs):
    """Injector flows, O/F, gas properties and nozzle flow at a given chamber pressure Pc.

    At the balanced Pc (from chamber_balance), injector flow in equals nozzle flow out.
    During the startup ramp Pc is held below that, so more flows in than out (the chamber filling).
    """
    ox = ox_flow(P_ox, Pc, up, ic, inputs)
    m_fuel = fuel_flow(P_fuel, Pc, ic, inputs)
    OF = ox["mdot"] / m_fuel if m_fuel > 0 else OF_MAX  # no fuel flow: let combustion clamp it
    gas = combustion(Pc, OF, cea, inputs)
    P_man_fuel = Pc + m_fuel ** 2 / (2 * ic["rho_fuel"] * ic["CdA_fuel"] ** 2)  # injector's share of the drop

    return {
        "Pc": Pc,
        "mdot_ox": ox["mdot"],
        "mdot_fuel": m_fuel,
        "mdot": ox["mdot"] + m_fuel,                 # injector flow in
        "mdot_nozzle": Pc * ic["A_t"] / gas["cstar"],  # nozzle flow out (c* = Pc * At / mdot, rearranged)
        "OF": OF,
        "P_man_ox": ox["P_man"],
        "P_man_fuel": P_man_fuel,
        "stiffness_ox": (ox["P_man"] - Pc) / Pc,
        "stiffness_fuel": (P_man_fuel - Pc) / Pc,
        **gas,  # cstar, Tc, MW, gamma, R
    }


def chamber_balance(P_ox, P_fuel, up, cea, ic, inputs, tol=100.0):
    """Chamber pressure where injector flow in equals nozzle flow out, plus everything at that point.

    P_ox, P_fuel: tank pressures this step [Pa]; up: from ox_upstream (once per step);
    cea: table from cea_table; tol: bisection tolerance on Pc [Pa].
    """
    def mismatch(Pc):
        """Injector flow in minus nozzle flow out."""
        cs = chamber_state(Pc, P_ox, P_fuel, up, cea, ic, inputs)
        return cs["mdot"] - cs["mdot_nozzle"]

    # At ambient the injectors flow most and the nozzle almost nothing (positive); at the fuel
    # tank pressure fuel stops and the nozzle flow is large (negative), so Pc lies in between
    Pc = bisect(mismatch, inputs["Ambient Pressure"], P_fuel, tol=tol)
    return chamber_state(Pc, P_ox, P_fuel, up, cea, ic, inputs)


def area_ratio(M, gamma):
    """Isentropic area ratio A/A_throat at Mach M (throat choked at Mach 1)."""
    return (1 / M) * ((2 / (gamma + 1)) * (1 + (gamma - 1) / 2 * M ** 2)) ** ((gamma + 1) / (2 * (gamma - 1)))


def nozzle(cb, ic, inputs):
    """Exit-plane conditions and thrust for the chamber state cb (from chamber_balance).

    Isentropic expansion with a choked throat and the flow attached to the wall; the
    separation flag marks when that last assumption becomes doubtful.
    """
    gamma, Pc = cb["gamma"], cb["Pc"]
    Pa = inputs["Ambient Pressure"]

    # Exit Mach: the area-Mach relation can't be solved for M, so bisect on it. Searching
    # from Mach 1 up picks the supersonic solution (each area ratio also has a subsonic one).
    Me = bisect(lambda M: area_ratio(M, gamma) - ic["eps"], 1.0, 10.0, tol=1e-6)

    # Isentropic relations: the gas cools and loses pressure as it speeds up
    temp_ratio = 1 + (gamma - 1) / 2 * Me ** 2
    Pe = Pc * temp_ratio ** (-gamma / (gamma - 1))
    Te = cb["Tc"] / temp_ratio

    # Exhaust velocity = Mach * local speed of sound, less nozzle losses
    ve = inputs["Nozzle Efficiency"] * Me * math.sqrt(gamma * cb["R"] * Te)

    # Thrust: momentum of the gas leaving the nozzle (with conical divergence loss) + pressure
    # on the exit area. Uses nozzle flow, so during the startup ramp thrust rises with Pc.
    # Floored at 0: at very low Pc the nozzle isn't flowing properly and the formula goes negative.
    divergence = (1 + math.cos(inputs["Exit Half Angle"])) / 2
    F_momentum = divergence * cb["mdot_nozzle"] * ve
    F_pressure = (Pe - Pa) * ic["A_e"]
    F = max(F_momentum + F_pressure, 0.0)

    return {
        "Me": Me,
        "Pe": Pe,
        "Te": Te,
        "ve": ve,
        "F_momentum": F_momentum,
        "F_pressure": F_pressure,
        "F": F,
        "Isp": F / (cb["mdot_nozzle"] * G0) if cb["mdot_nozzle"] > 0 else 0.0,
        "separation_risk": bool(Pe < SEPARATION_RATIO * Pa),  # bool(): CEA's numpy values give numpy booleans
    }


def motor_class(total_impulse):
    """Motor impulse class letter: A = 1.26-2.5 N*s, and each later letter doubles the range."""
    if total_impulse <= 2.5:
        return "A"
    return chr(ord("A") + math.ceil(math.log2(total_impulse / 2.5)))


def run(inputs, progress_every=100):
    """Simulate one burn. Returns {"history": arrays vs. time, "summary": totals and stop reason}."""
    ic = initial_conditions(inputs)
    cea = cea_table(inputs)
    dt = inputs["Timestep"]
    t_end = inputs["Sim Duration"]
    Pa = inputs["Ambient Pressure"]

    # State: the only quantities carried from one step to the next
    t = 0.0
    m_ox = ic["m_ox0"]
    m_fuel = ic["m_fuel0"]
    t_switch = None

    history = {}  # name -> list of values, one per step
    stop_reason = "sim duration reached"
    step = 0

    while t < t_end:
        tp = tank_pressures(m_ox, t, t_switch, ic, inputs)
        t_switch = tp["t_switch"]

        # Stop checks, before any physics that would fail in these conditions
        if m_ox <= 0:
            stop_reason = "oxidizer depleted"
            break
        if m_fuel <= 0:
            stop_reason = "fuel depleted"
            break
        if tp["P_fuel"] <= Pa:
            stop_reason = "fuel tank pressure at or below ambient"
            break

        # Physics for this instant, all from the current state
        up = ox_upstream(tp["P_ox"], inputs, tp["phase"])
        cb = chamber_balance(tp["P_ox"], tp["P_fuel"], up, cea, ic, inputs)
        if t < STARTUP_TIME:
            # Startup ramp: hold Pc between ambient and its balanced value, then re-evaluate there
            ramp = t / STARTUP_TIME
            cb = chamber_state(Pa + ramp * (cb["Pc"] - Pa), tp["P_ox"], tp["P_fuel"], up, cea, ic, inputs)
            step_dt = STARTUP_DT
        else:
            step_dt = dt
        nz = nozzle(cb, ic, inputs)

        # Record everything at time t
        record = {"t": t, "dt": step_dt, "m_ox": m_ox, "m_fuel": m_fuel,
                  "P_ox": tp["P_ox"], "P_fuel": tp["P_fuel"], **cb, **nz}
        for name, value in record.items():
            history.setdefault(name, []).append(value)

        # Step the state forward (Euler): drain the tanks at this instant's injector flow rates
        m_ox = max(m_ox - cb["mdot_ox"] * step_dt, 0.0)
        m_fuel = max(m_fuel - cb["mdot_fuel"] * step_dt, 0.0)
        t += step_dt

        step += 1
        if progress_every and step % progress_every == 0:
            # The end time isn't known in advance; progress = how close the first tank is to empty
            used = max(1 - m_ox / ic["m_ox0"], 1 - m_fuel / ic["m_fuel0"])
            print(f"  {used:4.0%} propellant used   t = {t:5.2f} s   F = {nz['F']:6.0f} N   "
                  f"Pc = {cb['Pc'] / 6894.757:5.1f} psi")

    history = {name: np.array(values) for name, values in history.items()}

    # Summary (each step's thrust acts for that step's dt, matching the Euler update)
    F = history.get("F", np.array([]))
    burn_time = float(np.sum(history["dt"])) if len(F) else 0.0
    total_impulse = float(np.sum(F * history["dt"])) if len(F) else 0.0
    prop_used = (ic["m_ox0"] - m_ox) + (ic["m_fuel0"] - m_fuel)
    summary = {
        "stop_reason": stop_reason,
        "ambient_pressure": Pa,
        "burn_time": burn_time,
        "total_impulse": total_impulse,
        "average_thrust": total_impulse / burn_time if burn_time > 0 else 0.0,
        "peak_thrust": float(F.max()) if len(F) else 0.0,
        "average_Isp": total_impulse / (prop_used * G0) if prop_used > 0 else 0.0,
        "ox_used": ic["m_ox0"] - m_ox,
        "fuel_used": ic["m_fuel0"] - m_fuel,
        "ox_left": m_ox,
        "fuel_left": m_fuel,
        "min_stiffness_ox": float(history["stiffness_ox"].min()) if len(F) else 0.0,
        "min_stiffness_fuel": float(history["stiffness_fuel"].min()) if len(F) else 0.0,
        # after the startup ramp (low Pc during the ramp always trips it)
        "separation_risk_any": bool(history["separation_risk"][history["t"] >= STARTUP_TIME].any()) if len(F) else False,
    }
    summary["motor"] = f"{motor_class(total_impulse)}{summary['average_thrust']:.0f}"

    return {"history": history, "summary": summary}


# ---------- sweeps ----------
def run_sweeps(inputs, sweeps):
    """Run the sim for every combination of the active sweeps (one sweep: a list; two: a grid).

    Each case starts from a fresh copy of the inputs with the swept parameters overridden.
    Returns a list of cases: {"values": [value per sweep, spreadsheet units], "result" or "error"}.
    Note: spreadsheet formulas (e.g. Oxidizer/Fuel Length, Nitrous Temperature) are not
    recalculated, so sweeping an input they depend on doesn't change them.
    """
    cases = []
    combos = list(itertools.product(*[range(len(s["values"])) for s in sweeps]))
    start = time.perf_counter()
    for n, idx in enumerate(combos, start=1):
        case_inputs = dict(inputs)
        for s, k in zip(sweeps, idx):
            case_inputs[s["parameter"]] = s["values"][k]
        values = [s["as_entered"][k] for s, k in zip(sweeps, idx)]
        label = ", ".join(f"{s['parameter']} = {v:g} {s['unit']}" for s, v in zip(sweeps, values))
        try:
            result = run(case_inputs, progress_every=0)
            sm = result["summary"]
            outcome = (f"{sm['motor']}, I = {sm['total_impulse']:.0f} N*s, "
                       f"Isp = {sm['average_Isp']:.0f} s ({sm['stop_reason']})")
            cases.append({"index": idx, "values": values, "result": result})
        except Exception as e:  # an extreme design shouldn't stop the whole sweep
            outcome = f"FAILED ({e})"
            cases.append({"index": idx, "values": values, "result": None, "error": str(e)})

        elapsed = time.perf_counter() - start
        remaining = elapsed / n * (len(combos) - n)
        print(f"[{n}/{len(combos)} {n / len(combos):4.0%}, ~{format_duration(remaining)} left] {label}: {outcome}")
    print(f"Sweep done in {format_duration(time.perf_counter() - start)}")
    return cases


def format_duration(seconds):
    """e.g. 75 -> '1m15s', 8 -> '8s'."""
    seconds = round(seconds)
    return f"{seconds // 60}m{seconds % 60:02d}s" if seconds >= 60 else f"{seconds}s"


SWEEP_METRICS = [  # (summary key, label, scale)
    ("total_impulse", "Total impulse [N·s]", 1),
    ("average_thrust", "Average thrust [N]", 1),
    ("peak_thrust", "Peak thrust [N]", 1),
    ("average_Isp", "Average Isp [s]", 1),
    ("burn_time", "Burn time [s]", 1),
    ("min_stiffness_ox", "Min oxidizer stiffness [%]", 100),
    ("min_stiffness_fuel", "Min fuel stiffness [%]", 100),
]


def save_sweep_csv(sweeps, cases, path="outputs/sweep_results.csv"):
    """One row per case: swept values (spreadsheet units) and the summary numbers."""
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow([f"{s['parameter']} [{s['unit']}]" for s in sweeps]
                   + ["motor", "stop reason"] + [label for _, label, _ in SWEEP_METRICS])
        for c in cases:
            if c["result"] is None:
                w.writerow(c["values"] + ["FAILED", c["error"]])
                continue
            sm = c["result"]["summary"]
            w.writerow(c["values"] + [sm["motor"], sm["stop_reason"]]
                       + [f"{sm[key] * scale:.6g}" for key, _, scale in SWEEP_METRICS])


# ---------- plotting (interactive web pages: hover for values, drag to zoom, double-click to reset) ----------
PSI = 6894.757  # Pa per psi

THEME = "dark"          # "dark" or "light"
OPEN_IN_BROWSER = True   # open each page when it's written

# Colorblind-validated reference palette, stepped separately for each background.
# categorical: fixed order, so series keep their color across plots.
# sequential: for ordered values (sweep cases, heat maps), from "low" (recedes into the
#   background) to "high" (stands out): light -> dark blue on light, dark -> light blue on dark.
THEMES = {
    "light": dict(
        surface="#fcfcfb", text="#0b0b0b", text_secondary="#52514e", grid="#e6e5e1", guide="#8a8a86",
        categorical=["#2a78d6", "#eb6834", "#1baf7a", "#eda100", "#e87ba4", "#008300", "#6250d6", "#e34948"],
        sequential=["#cde2fb", "#86b6ef", "#2a78d6", "#1c5cab", "#104281"],
        legend_bg="rgba(252,252,251,0.85)",
    ),
    "dark": dict(
        surface="#1a1a19", text="#f0efec", text_secondary="#c3c2b7", grid="#383835", guide="#8a8a86",
        categorical=["#3987e5", "#d95926", "#199e70", "#c98500", "#d55181", "#008300", "#9085e9", "#e66767"],
        sequential=["#0d366b", "#184f95", "#2a78d6", "#86b6ef", "#cde2fb"],
        legend_bg="rgba(26,26,25,0.85)",
    ),
}
STYLE = THEMES[THEME]
CATEGORICAL = STYLE["categorical"]
SEQUENTIAL = STYLE["sequential"]
GUIDE_COLOR = STYLE["guide"]  # reference lines (ambient, guidelines)


def sequential_colors(n):
    """n colors along the sequential ramp, for ordered values such as a sweep (skips the step that
    would blend into the background)."""
    stops = [tuple(int(c[i:i + 2], 16) for i in (1, 3, 5)) for c in SEQUENTIAL]
    colors = []
    for k in range(n):
        x = (0.3 + 0.7 * k / max(n - 1, 1)) * (len(stops) - 1)
        i = min(int(x), len(stops) - 2)
        f = x - i
        rgb = [round(a + (b - a) * f) for a, b in zip(stops[i], stops[i + 1])]
        colors.append("#%02x%02x%02x" % tuple(rgb))
    return colors


def save_page(fig, path, title):
    """Write the figure as a standalone web page (works offline) and optionally open it."""
    s = STYLE
    fig.update_layout(title=title, template="plotly_dark" if THEME == "dark" else "plotly_white",
                      paper_bgcolor=s["surface"], plot_bgcolor=s["surface"],
                      font=dict(size=12, color=s["text"]), margin=dict(l=60, r=30, t=90, b=40))
    fig.update_xaxes(gridcolor=s["grid"], zerolinecolor=s["grid"], linecolor=s["grid"])
    fig.update_yaxes(gridcolor=s["grid"], zerolinecolor=s["grid"], linecolor=s["grid"])
    fig.update_annotations(font_color=s["text_secondary"])  # subplot titles and reference-line labels
    os.makedirs(os.path.dirname(path), exist_ok=True)
    html = fig.to_html(include_plotlyjs=True, full_html=True)
    # Match the page itself to the plot background, so there's no white border around it
    html = html.replace("<head>", f"<head><style>body {{ background: {s['surface']}; margin: 0; }}</style>", 1)
    with open(path, "w", encoding="utf-8") as f:
        f.write(html)
    if OPEN_IN_BROWSER:
        webbrowser.open(Path(path).resolve().as_uri())
    return path


def panel_legend(fig, n, row, col, title=None, corner="right"):
    """Give one subplot its own legend (the n-th, counting from 1) in its top-right (or top-left)
    corner; returns the legend id to put on its traces."""
    lid = "legend" if n == 1 else f"legend{n}"
    sp = fig.get_subplot(row, col)
    x = sp.xaxis.domain[1] if corner == "right" else sp.xaxis.domain[0] + 0.01
    fig.update_layout(**{lid: dict(x=x, y=sp.yaxis.domain[1], xanchor=corner, yanchor="top",
                                   bgcolor=STYLE["legend_bg"], font=dict(size=10),
                                   title=dict(text=title) if title else None)})
    return lid


def line(fig, row, col, x, y, name, color, dash=None, hover_unit="", legend=None):
    """A line on one subplot; it appears in that subplot's legend if one is given."""
    fig.add_trace(go.Scatter(x=x, y=y, name=name, mode="lines", line=dict(color=color, width=2, dash=dash),
                             showlegend=legend is not None, legend=legend or "legend",
                             hovertemplate=f"{name}: %{{y:.4g}} {hover_unit}<extra></extra>"), row=row, col=col)


def plot_results(result, path="outputs/results.html"):
    """Single run: thrust, pressures, stiffness, flows, O/F, Isp, propellant, c*, Tc, exit pressure vs. time."""
    h, s = result["history"], result["summary"]
    t = h["t"]
    burn = t >= STARTUP_TIME  # during the ramp Pc is near ambient, so ratios like stiffness spike meaninglessly
    tb = t[burn]
    blue, orange, aqua, yellow, magenta = CATEGORICAL[:5]

    titles = ["Thrust [N]", "Pressure [psi]", "Injector stiffness [%]", "Mass flow [kg/s]", "O/F",
              "Isp [s]", "Propellant left [kg]", "c* with efficiency [m/s]", "Chamber temperature [K]",
              "Nozzle exit pressure [psi]"]
    fig = make_subplots(rows=5, cols=2, shared_xaxes=True, subplot_titles=titles,
                        vertical_spacing=0.05, horizontal_spacing=0.08)

    lg = panel_legend(fig, 1, 1, 1)
    line(fig, 1, 1, t, h["F"], "total", blue, hover_unit="N", legend=lg)
    line(fig, 1, 1, t, h["F_momentum"], "momentum", orange, "dash", "N", legend=lg)
    line(fig, 1, 1, t, h["F_pressure"], "pressure", aqua, "dot", "N", legend=lg)

    lg = panel_legend(fig, 2, 1, 2)
    for (key, name), color in zip([("P_ox", "ox tank"), ("P_fuel", "fuel tank"), ("P_man_ox", "ox manifold"),
                                   ("P_man_fuel", "fuel manifold"), ("Pc", "chamber")], CATEGORICAL):
        line(fig, 1, 2, t, h[key] / PSI, name, color, hover_unit="psi", legend=lg)

    lg = panel_legend(fig, 3, 2, 1)
    line(fig, 2, 1, tb, h["stiffness_ox"][burn] * 100, "oxidizer", blue, hover_unit="%", legend=lg)
    line(fig, 2, 1, tb, h["stiffness_fuel"][burn] * 100, "fuel", orange, hover_unit="%", legend=lg)
    fig.add_hline(y=30, line=dict(color=GUIDE_COLOR, dash="dot"), annotation_text="30% guideline", row=2, col=1)

    lg = panel_legend(fig, 4, 2, 2)
    line(fig, 2, 2, t, h["mdot_ox"], "oxidizer", blue, hover_unit="kg/s", legend=lg)
    line(fig, 2, 2, t, h["mdot_fuel"], "fuel", orange, hover_unit="kg/s", legend=lg)
    line(fig, 3, 1, tb, h["OF"][burn], "O/F", blue)
    line(fig, 3, 2, tb, h["Isp"][burn], "Isp", blue, hover_unit="s")
    lg = panel_legend(fig, 5, 4, 1)
    line(fig, 4, 1, t, h["m_ox"], "oxidizer", blue, hover_unit="kg", legend=lg)
    line(fig, 4, 1, t, h["m_fuel"], "fuel", orange, hover_unit="kg", legend=lg)
    line(fig, 4, 2, tb, h["cstar"][burn], "c*", blue, hover_unit="m/s")
    line(fig, 5, 1, tb, h["Tc"][burn], "chamber temperature", blue, hover_unit="K")
    line(fig, 5, 2, tb, h["Pe"][burn] / PSI, "exit pressure", blue, hover_unit="psi")
    Pa = s["ambient_pressure"] / PSI
    fig.add_hline(y=Pa, line=dict(color=GUIDE_COLOR, dash="dash"), annotation_text="ambient", row=5, col=2)
    fig.add_hline(y=SEPARATION_RATIO * Pa, line=dict(color=GUIDE_COLOR, dash="dot"),
                  annotation_text="separation risk", row=5, col=2)

    for col in (1, 2):
        fig.update_xaxes(title_text="Time [s]", row=5, col=col)
    fig.update_layout(height=1600, hovermode="x unified")
    title = (f"{s['motor']}   I = {s['total_impulse']:.0f} N·s   burn = {s['burn_time']:.2f} s   "
             f"Isp = {s['average_Isp']:.0f} s   (stop: {s['stop_reason']})")
    return save_page(fig, path, title)


def plot_sweep(sweeps, cases, path="outputs/sweep.html"):
    """One sweep: curves over time for every value, plus summary metrics vs. the value.
    Two sweeps: a heat map of each summary metric over the grid."""
    if len(sweeps) == 1:
        return _plot_sweep_1d(sweeps[0], cases, path)
    return _plot_sweep_2d(sweeps, cases, path)


def _plot_sweep_1d(sweep, cases, path, max_legend=12):
    ok = [c for c in cases if c["result"] is not None]
    colors = sequential_colors(len(ok))
    unit = f" [{sweep['unit']}]" if sweep["unit"] else ""
    label = f"{sweep['parameter']}{unit}"
    x = [c["values"][0] for c in ok]

    curves = [("F", "Thrust [N]", 1, False), ("Pc", "Chamber pressure [psi]", 1 / PSI, False),
              ("OF", "O/F", 1, True), ("stiffness_ox", "Oxidizer injector stiffness [%]", 100, True),
              ("stiffness_fuel", "Fuel injector stiffness [%]", 100, True)]
    metrics = [  # (title, [(summary key, scale, series name, color)])
        ("Total impulse [N·s]", [("total_impulse", 1, "total impulse", CATEGORICAL[0])]),
        ("Average Isp [s]", [("average_Isp", 1, "average Isp", CATEGORICAL[0])]),
        ("Thrust [N]", [("average_thrust", 1, "average thrust", CATEGORICAL[0]),
                        ("peak_thrust", 1, "peak thrust", CATEGORICAL[1])]),
        ("Burn time [s]", [("burn_time", 1, "burn time", CATEGORICAL[0])]),
        ("Minimum injector stiffness [%]", [("min_stiffness_ox", 100, "min ox stiffness", CATEGORICAL[0]),
                                            ("min_stiffness_fuel", 100, "min fuel stiffness", CATEGORICAL[1])]),
    ]
    titles = [t for pair in zip([c[1] + " vs. time" for c in curves], [m[0] for m in metrics]) for t in pair]
    fig = make_subplots(rows=5, cols=2, subplot_titles=titles, vertical_spacing=0.05, horizontal_spacing=0.08)

    # Case values get a legend in the thrust panel (while there are few enough to list); clicking an
    # entry toggles that case in every panel, since its curves share a legend group
    n_legend = 1
    values_legend = panel_legend(fig, n_legend, 1, 1, title=label) if len(ok) <= max_legend else None
    for row, (key, title, scale, burn_only) in enumerate(curves, start=1):
        for c, color in zip(ok, colors):
            h = c["result"]["history"]
            keep = h["t"] >= STARTUP_TIME if burn_only else np.ones_like(h["t"], bool)
            v = c["values"][0]
            fig.add_trace(go.Scatter(
                x=h["t"][keep], y=h[key][keep] * scale, mode="lines", line=dict(color=color, width=1.5),
                name=f"{v:g}", legendgroup=f"{v:g}", legend=values_legend or "legend",
                showlegend=(row == 1 and values_legend is not None),
                hovertemplate=f"{label} = {v:g}<br>t = %{{x:.2f}} s<br>%{{y:.4g}}<extra></extra>"), row=row, col=1)
        fig.update_xaxes(title_text="Time [s]", row=row, col=1)
    fig.add_hline(y=30, line=dict(color=GUIDE_COLOR, dash="dot"), annotation_text="30% guideline", row=4, col=1)
    fig.add_hline(y=30, line=dict(color=GUIDE_COLOR, dash="dot"), annotation_text="30% guideline", row=5, col=1)

    for row, (title, series) in enumerate(metrics, start=1):
        lg = None
        if len(series) > 1:
            n_legend += 1
            lg = panel_legend(fig, n_legend, row, 2, corner="left")
        for key, scale, name, color in series:
            fig.add_trace(go.Scatter(
                x=x, y=[c["result"]["summary"][key] * scale for c in ok], mode="lines+markers",
                line=dict(color=color, width=2), marker=dict(size=7), name=name,
                showlegend=lg is not None, legend=lg or "legend",
                hovertemplate=f"{label} = %{{x:g}}<br>{name} = %{{y:.4g}}<extra></extra>"), row=row, col=2)
        fig.update_xaxes(title_text=label, row=row, col=2)
    fig.add_hline(y=30, line=dict(color=GUIDE_COLOR, dash="dot"), annotation_text="30% guideline", row=5, col=2)

    fig.update_layout(height=1700, hovermode="closest")
    note = "" if len(ok) <= max_legend else "   (hover a curve to see its value)"
    return save_page(fig, path, f"Sweep: {label}, {len(ok)} cases{note}")


def _plot_sweep_2d(sweeps, cases, path):
    s1, s2 = sweeps
    label1 = f"{s1['parameter']}" + (f" [{s1['unit']}]" if s1["unit"] else "")
    label2 = f"{s2['parameter']}" + (f" [{s2['unit']}]" if s2["unit"] else "")
    shown = [SWEEP_METRICS[i] for i in (0, 3, 2, 4, 5, 6)]  # impulse, Isp, peak thrust, burn time, min stiffness x2

    fig = make_subplots(rows=3, cols=2, subplot_titles=[m[1] for m in shown],
                        vertical_spacing=0.1, horizontal_spacing=0.14)
    small = len(s1["values"]) * len(s2["values"]) <= 64  # print values in the cells only while they're readable
    for n, (key, title, scale) in enumerate(shown):
        row, col = n // 2 + 1, n % 2 + 1
        grid = np.full((len(s2["values"]), len(s1["values"])), np.nan)  # rows: sweep 2, columns: sweep 1
        for c in cases:
            if c["result"] is not None:
                i1, i2 = c["index"]
                grid[i2, i1] = c["result"]["summary"][key] * scale
        fig.add_trace(go.Heatmap(
            x=[f"{v:g}" for v in s1["as_entered"]], y=[f"{v:g}" for v in s2["as_entered"]], z=grid,
            colorscale=[[i / (len(SEQUENTIAL) - 1), c] for i, c in enumerate(SEQUENTIAL)],
            colorbar=dict(len=0.27, y=1 - (row - 0.5) / 3, x=0.43 if col == 1 else 1.0, thickness=12),
            texttemplate="%{z:.4g}" if small else None,
            hovertemplate=f"{label1} = %{{x}}<br>{label2} = %{{y}}<br>{title} = %{{z:.4g}}<extra></extra>"),
            row=row, col=col)
        fig.update_xaxes(title_text=label1, row=row, col=col)
        fig.update_yaxes(title_text=label2, row=row, col=col)

    fig.update_layout(height=1300)
    return save_page(fig, path, f"Sweep: {s1['parameter']} × {s2['parameter']}")


if __name__ == "__main__":
    inputs, sweeps = load_inputs()

    if sweeps:
        n_cases = math.prod(len(s["values"]) for s in sweeps)
        print(f"Running {n_cases} cases: " + " x ".join(f"{s['parameter']} ({len(s['values'])})" for s in sweeps))
        cases = run_sweeps(inputs, sweeps)
        save_sweep_csv(sweeps, cases)
        plot_sweep(sweeps, cases)
        print("\nSaved outputs/sweep.html and outputs/sweep_results.csv")
    else:
        print("Running...")
        result = run(inputs)
        print()
        for name, value in result["summary"].items():
            print(f"{name:22} {value:.6g}" if isinstance(value, float) else f"{name:22} {value}")
        plot_results(result)
        print("\nSaved outputs/results.html")
